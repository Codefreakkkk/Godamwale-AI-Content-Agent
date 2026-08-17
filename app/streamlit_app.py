import streamlit as st

from docx import Document

from agent_controller import run_agent
from website_reader import extract_website_content
from content_quality import calculate_content_quality

from config import (
    MAX_WEBSITES
)


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Godamwale AI Content Agent",
    page_icon="🤖",
    layout="wide"
)


# ============================================================
# SIDEBAR
# ============================================================

st.sidebar.title(
    "🤖 Godamwale AI Content Agent"
)

st.sidebar.write(
    "AI-powered content assistant for:"
)

st.sidebar.write(
    """
    ✅ Website Analysis

    ✅ Website Content Optimization

    ✅ SEO & AIO Analysis

    ✅ Blog Generation + Optimization

    ✅ Competitor Analysis
    """
)

st.sidebar.divider()

st.sidebar.subheader(
    "🌐 Website Analysis"
)

st.sidebar.write(
    f"""
    🌐 Analyze up to {MAX_WEBSITES} websites

    📄 Extract website content

    📊 View page and token statistics

    🎯 Analyze content around target keywords

    ⚔️ Compare competitor content
    """
)

st.sidebar.divider()

st.sidebar.subheader(
    "⚙️ System Status"
)

st.sidebar.write(
    """
    🟢 Website Content Extraction

    🟢 AI Content Modules

    🟢 Content Quality Analysis

    🟢 RAG Foundation

    🟡 AI Provider Status Depends On API
    """
)

st.sidebar.divider()

st.sidebar.caption(
    "Godamwale AI Content Agent"
)


# ============================================================
# MAIN PAGE
# ============================================================

st.title(
    "🤖 Godamwale AI Content Agent"
)

st.caption(
    "Analyze, optimize, compare, and generate "
    "business content using AI."
)


# ============================================================
# WEBSITE ANALYSIS
# ============================================================

st.subheader(
    "🌐 Website Analysis"
)

st.write(
    f"Enter up to {MAX_WEBSITES} website URLs for analysis."
)

website_urls_text = st.text_area(
    "Website URLs",
    placeholder=(
        f"Enter one website URL per line.\n\n"
        f"Maximum: {MAX_WEBSITES} websites\n\n"
        "Example:\n"
        "https://www.godamwale.com/\n"
        "https://example.com"
    ),
    height=120
)

target_keyword = st.text_input(
    "🎯 Target Keyword",
    placeholder="Example: warehouse in Kolkata"
)


# ============================================================
# LOAD WEBSITES
# ============================================================

if st.button(
    "🌐 Load Websites"
):

    if not website_urls_text.strip():

        st.error(
            "Please enter at least one website URL."
        )

    else:

        website_urls = [
            url.strip()
            for url in website_urls_text.splitlines()
            if url.strip()
        ]

        if len(website_urls) > MAX_WEBSITES:

            st.error(
                f"Please enter a maximum of "
                f"{MAX_WEBSITES} website URLs."
            )

        else:

            combined_website_content = ""

            all_statistics = []

            successful_websites = 0

            with st.spinner(
                "Reading website content..."
            ):

                for website_url in website_urls:

                    try:

                        (
                            extracted_content,
                            statistics
                        ) = extract_website_content(
                            website_url
                        )

                        if extracted_content:

                            combined_website_content += (
                                "\n\n"
                                "========================================\n"
                                f"WEBSITE: {website_url}\n"
                                "========================================\n\n"
                                + extracted_content
                            )

                            all_statistics.append(
                                {
                                    "website": website_url,
                                    "statistics": statistics
                                }
                            )

                            successful_websites += 1

                        else:

                            st.warning(
                                f"Unable to extract content from: "
                                f"{website_url}"
                            )

                    except Exception as error:

                        st.warning(
                            f"Unable to process: "
                            f"{website_url}"
                        )

                        print(
                            "Website Error:",
                            error
                        )


            # =================================================
            # STORE WEBSITE DATA
            # =================================================

            if successful_websites > 0:

                st.session_state[
                    "website_content"
                ] = combined_website_content

                st.session_state[
                    "website_urls"
                ] = website_urls

                st.session_state[
                    "website_statistics"
                ] = all_statistics

                st.success(
                    f"{successful_websites} website(s) "
                    f"loaded successfully."
                )

                st.info(
                    "Website content is now available "
                    "to the AI Agent."
                )


                # =============================================
                # WEBSITE STATISTICS
                # =============================================

                st.subheader(
                    "📊 Website Content Analysis"
                )

                for website_data in all_statistics:

                    website = website_data[
                        "website"
                    ]

                    statistics = website_data[
                        "statistics"
                    ]

                    st.write(
                        f"### 🌐 {website}"
                    )

                    col1, col2, col3 = st.columns(
                        3
                    )

                    with col1:

                        st.metric(
                            "Pages",
                            statistics[
                                "total_pages"
                            ]
                        )

                    with col2:

                        st.metric(
                            "Characters",
                            statistics[
                                "total_characters"
                            ]
                        )

                    with col3:

                        st.metric(
                            "Estimated Tokens",
                            statistics[
                                "total_estimated_tokens"
                            ]
                        )

                    if statistics[
                        "pages"
                    ]:

                        st.write(
                            "**Page-level breakdown:**"
                        )

                        for page in statistics[
                            "pages"
                        ]:

                            st.write(
                                f"- `{page['url']}` — "
                                f"{page['characters']} characters, "
                                f"~{page['estimated_tokens']} tokens"
                            )

            else:

                st.error(
                    "No website content could be extracted."
                )


# ============================================================
# SEPARATOR
# ============================================================

st.divider()


# ============================================================
# TASK SELECTION
# ============================================================

st.subheader(
    "🧠 Choose Content Task"
)

task_option = st.selectbox(
    "Select task:",
    [
        "Website Content Optimization",
        "SEO & AIO Analysis",
        "Blog Generation + Optimization",
        "Competitor Analysis"
    ]
)


# ============================================================
# DIRECT MODULE MAPPING
# ============================================================

TASK_TO_MODULE = {

    "Website Content Optimization":
        "optimization",

    "SEO & AIO Analysis":
        "seo",

    "Blog Generation + Optimization":
        "blog",

    "Competitor Analysis":
        "competitor_analysis"
}


selected_module = TASK_TO_MODULE[
    task_option
]


# ============================================================
# TASK DESCRIPTION
# ============================================================

if task_option == "Website Content Optimization":

    st.caption(
        "Improve existing website content for clarity, "
        "structure, and business relevance."
    )

elif task_option == "SEO & AIO Analysis":

    st.caption(
        "Analyze and improve content for search and "
        "AI-oriented visibility."
    )

elif task_option == "Blog Generation + Optimization":

    st.caption(
        "Generate a blog from approved information, "
        "then optimize the generated blog for the target keyword."
    )

elif task_option == "Competitor Analysis":

    st.caption(
        "Compare Godamwale content with competitor content "
        "and identify practical content gaps."
    )


# ============================================================
# COMPETITOR WEBSITE INPUT
# ============================================================

competitor_url = ""

if task_option == "Competitor Analysis":

    st.subheader(
        "⚔️ Competitor Website"
    )

    competitor_url = st.text_input(
        "Competitor Website URL",
        placeholder="Example: https://competitor.com"
    )


# ============================================================
# ADDITIONAL INSTRUCTIONS
# ============================================================

additional_instructions = st.text_area(
    "📝 Additional Instructions",
    placeholder=(
        "Example:\n"
        "Target Indian businesses.\n"
        "Keep the content professional.\n"
        "Focus on warehouse services."
    ),
    height=120
)


# ============================================================
# BUILD USER REQUEST
# ============================================================

if task_option == "Blog Generation + Optimization":

    user_request = "Blog Generation"

else:

    user_request = task_option


if target_keyword:

    user_request += (
        "\n\nTarget Keyword:\n"
        + target_keyword
    )


if additional_instructions:

    user_request += (
        "\n\nAdditional Instructions:\n"
        + additional_instructions
    )


# ============================================================
# RUN AI AGENT
# ============================================================

if st.button(
    "🚀 Run AI Agent"
):

    # --------------------------------------------------------
    # Validate website
    # --------------------------------------------------------

    if not st.session_state.get(
        "website_content"
    ):

        st.warning(
            "Please load at least one website before "
            "running the AI Agent."
        )

    # --------------------------------------------------------
    # Validate competitor URL
    # --------------------------------------------------------

    elif (
        task_option == "Competitor Analysis"
        and not competitor_url.strip()
    ):

        st.warning(
            "Please enter a competitor website URL."
        )

    else:

        competitor_content = None


        # ====================================================
        # LOAD COMPETITOR WEBSITE
        # ====================================================

        if task_option == "Competitor Analysis":

            with st.spinner(
                "Reading competitor website..."
            ):

                try:

                    (
                        competitor_content,
                        competitor_statistics
                    ) = extract_website_content(
                        competitor_url
                    )


                    if not competitor_content:

                        st.error(
                            "Unable to extract competitor "
                            "website content."
                        )

                        st.stop()


                    st.success(
                        "Competitor website loaded successfully."
                    )


                    # ----------------------------------------
                    # Competitor statistics
                    # ----------------------------------------

                    st.subheader(
                        "📊 Competitor Content Analysis"
                    )

                    col1, col2, col3 = st.columns(
                        3
                    )

                    with col1:

                        st.metric(
                            "Pages",
                            competitor_statistics[
                                "total_pages"
                            ]
                        )

                    with col2:

                        st.metric(
                            "Characters",
                            competitor_statistics[
                                "total_characters"
                            ]
                        )

                    with col3:

                        st.metric(
                            "Estimated Tokens",
                            competitor_statistics[
                                "total_estimated_tokens"
                            ]
                        )


                except Exception as error:

                    st.error(
                        "Unable to process competitor website."
                    )

                    print(
                        "Competitor Website Error:",
                        error
                    )

                    st.stop()


        # ====================================================
        # RUN AGENT
        # ====================================================

        with st.spinner(
            "AI Agent is working..."
        ):

            try:

                (
                    selected_module,
                    result
                ) = run_agent(
                    user_request=user_request,
                    website_content=st.session_state.get(
                        "website_content"
                    ),
                    competitor_content=competitor_content,
                    selected_module=selected_module
                )


                st.divider()


                # =================================================
                # SELECTED MODULE
                # =================================================

                st.subheader(
                    "🎯 Selected Module"
                )

                st.info(
                    selected_module.replace(
                        "_",
                        " "
                    ).title()
                )


                # =================================================
                # BLOG PIPELINE
                # =================================================

                if selected_module == "blog":

                    generated_blog = result[
                        "generated_blog"
                    ]

                    optimized_blog = result[
                        "optimized_blog"
                    ]


                    st.subheader(
                        "📝 Blog Generation → Optimization Pipeline"
                    )

                    st.caption(
                        "The AI Agent completed two stages: "
                        "generation followed by optimization."
                    )


                    # ---------------------------------------------
                    # STAGE 1
                    # ---------------------------------------------

                    st.markdown(
                        "### 📝 Stage 1 — Blog Generation"
                    )

                    st.info(
                        "The AI first generated the initial blog "
                        "using approved company information."
                    )

                    with st.expander(
                        "View Generated Blog",
                        expanded=True
                    ):

                        st.write(
                            generated_blog
                        )


                    # ---------------------------------------------
                    # STAGE 2
                    # ---------------------------------------------

                    st.divider()

                    st.markdown(
                        "### ✨ Stage 2 — Blog Optimization"
                    )

                    st.info(
                        "The generated blog was then optimized "
                        "for the selected target keyword."
                    )

                    with st.expander(
                        "View Optimized Blog",
                        expanded=True
                    ):

                        st.write(
                            optimized_blog
                        )


                    final_content = optimized_blog


                else:

                    # =================================================
                    # NORMAL AI RESPONSE
                    # =================================================

                    st.subheader(
                        "🤖 AI Generated Response"
                    )

                    ai_failed = (
                        isinstance(result, str)
                        and result.startswith(
                            "Unable to generate"
                        )
                    )


                    if ai_failed:

                        st.warning(
                            """
                            ⚠️ AI Response Unavailable

                            The AI generation service is currently unavailable.

                            Possible reasons:

                            - API usage limit reached
                            - Temporary AI provider issue
                            - Network problem

                            Please try again later.
                            """
                        )

                        final_content = ""


                    else:

                        st.write(
                            result
                        )

                        final_content = result


                # =================================================
                # CONTENT QUALITY
                # =================================================

                if final_content:

                    quality_report = (
                        calculate_content_quality(
                            final_content
                        )
                    )


                    st.divider()


                    st.subheader(
                        "📊 Content Quality Score"
                    )


                    st.metric(
                        "Overall Score",
                        f"{quality_report['score']} / 100"
                    )


                    st.subheader(
                        "Quality Breakdown"
                    )


                    for key, value in (
                        quality_report[
                            "details"
                        ].items()
                    ):

                        st.write(
                            f"**{key}:** {value}"
                        )


                    st.subheader(
                        "💡 Improvement Suggestions"
                    )


                    for suggestion in (
                        quality_report[
                            "suggestions"
                        ]
                    ):

                        st.write(
                            "• " + suggestion
                        )


                    # =================================================
                    # DOWNLOAD OPTIONS
                    # =================================================

                    st.divider()

                    st.subheader(
                        "📥 Download Final Content"
                    )


                    # -----------------------------------------------
                    # TXT
                    # -----------------------------------------------

                    st.download_button(
                        label="⬇️ Download as TXT",
                        data=final_content,
                        file_name="godamwale_ai_response.txt",
                        mime="text/plain"
                    )


                    # -----------------------------------------------
                    # DOCX
                    # -----------------------------------------------

                    doc = Document()

                    doc.add_heading(
                        "Godamwale AI Content Agent Response",
                        level=1
                    )

                    doc.add_paragraph(
                        final_content
                    )


                    docx_path = (
                        "godamwale_ai_response.docx"
                    )


                    doc.save(
                        docx_path
                    )


                    with open(
                        docx_path,
                        "rb"
                    ) as file:

                        st.download_button(
                            label="📄 Download as DOCX",
                            data=file,
                            file_name=docx_path,
                            mime=(
                                "application/vnd.openxmlformats-officedocument."
                                "wordprocessingml.document"
                            )
                        )


            except Exception as error:

                st.error(
                    "Unable to process request."
                )

                print(
                    "Agent Error:",
                    error
                )